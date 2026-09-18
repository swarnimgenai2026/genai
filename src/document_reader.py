"""Document ingestion: discover files in data/ and extract raw text.

Supports .txt, .pdf, and .docx. Each file is handled independently so a
single corrupt or unreadable document never aborts a batch run.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import List

from src.config import DATA_DIR, SUPPORTED_EXTENSIONS
from src.utils.logger import get_logger

logger = get_logger(__name__)


class DocumentReadError(Exception):
    """Raised when a file cannot be turned into usable text."""


@dataclass
class RawDocument:
    file_path: Path
    file_name: str
    text: str


def discover_documents(data_dir: Path = DATA_DIR) -> List[Path]:
    """Return eligible file paths in data_dir, sorted for reproducibility."""
    if not data_dir.exists():
        logger.warning("Data directory does not exist: %s", data_dir)
        return []

    files = [
        p for p in sorted(data_dir.iterdir())
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    logger.info("Discovered %d document(s) in %s", len(files), data_dir)
    return files


def _read_txt(path: Path) -> str:
    # Try encodings in order of likelihood; latin-1 never raises, so it
    # acts as a last-resort fallback.
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentReadError(f"Could not decode text file: {path.name}")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentReadError(
            "pypdf is required to read PDF files."
        ) from e

    try:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if not text:
            raise DocumentReadError(
                f"No extractable text in PDF (likely scanned/image-only): {path.name}"
            )
        return text
    except DocumentReadError:
        raise
    except Exception as e:
        raise DocumentReadError(f"Failed to parse PDF {path.name}: {e}") from e


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise DocumentReadError(
            "python-docx is required to read DOCX files."
        ) from e

    try:
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        # Complaint forms are sometimes laid out in tables rather than
        # plain paragraphs; capture those too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        text = "\n".join(paragraphs).strip()
        if not text:
            raise DocumentReadError(f"DOCX file appears to be empty: {path.name}")
        return text
    except DocumentReadError:
        raise
    except Exception as e:
        raise DocumentReadError(f"Failed to parse DOCX {path.name}: {e}") from e


_READERS = {
    ".txt": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}


def read_document(path: Path) -> RawDocument:
    """Extract text from one document, raising DocumentReadError on failure."""
    suffix = path.suffix.lower()
    reader_fn = _READERS.get(suffix)
    if reader_fn is None:
        raise DocumentReadError(f"Unsupported file type: {suffix}")

    if not path.exists():
        raise DocumentReadError(f"File does not exist: {path}")

    if path.stat().st_size == 0:
        raise DocumentReadError(f"File is empty: {path.name}")

    text = reader_fn(path).strip()
    if not text:
        raise DocumentReadError(f"No text extracted from: {path.name}")

    logger.debug("Read '%s' -- %d characters extracted.", path.name, len(text))
    return RawDocument(file_path=path, file_name=path.name, text=text)


def read_all_documents(data_dir: Path = DATA_DIR) -> List[RawDocument]:
    """Discover and read every eligible document; unreadable files are skipped.

    workflow.run_batch() does not use this directly -- it reads files
    itself so it can record failures as rows in the final report instead
    of silently dropping them.
    """
    documents = []
    for path in discover_documents(data_dir):
        try:
            documents.append(read_document(path))
        except DocumentReadError as e:
            logger.error("Skipping file due to read error: %s", e)
    return documents
